import io
from collections.abc import Iterator

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from src.ingestion.schemas import ParsedDocument, ParsedSection

_HEADING_STYLE_PREFIX = "Heading "
_TITLE_STYLE = "Title"


def _iter_block_items(document: DocumentObject) -> Iterator[Paragraph | Table]:
    """Yield paragraphs and tables in document order.

    `python-docx` exposes `document.paragraphs` and `document.tables` as
    separate, order-losing lists; this walks the underlying XML body
    instead so a table appearing between two paragraphs stays between them.
    """
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _heading_level(paragraph: Paragraph) -> int | None:
    # `.style.name` is `str | None` even for a real style (some are unnamed).
    style_name = (paragraph.style.name if paragraph.style else None) or ""
    if style_name == _TITLE_STYLE:
        return 0
    if style_name.startswith(_HEADING_STYLE_PREFIX):
        suffix = style_name[len(_HEADING_STYLE_PREFIX) :].strip()
        if suffix.isdigit():
            return int(suffix)
    return None


def _table_to_text(table: Table) -> str:
    rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
    return "\n".join(row for row in rows if row.strip())


def parse(content: bytes) -> ParsedDocument:
    """Parse a DOCX into a heading tree from its `Heading N` / `Title` styles.

    Body paragraphs and tables attach as text to the innermost currently
    open section; text before the first heading becomes a heading-less
    root section.
    """
    document = Document(io.BytesIO(content))

    root = ParsedSection(heading=None, level=0, text="", children=[])
    stack: list[ParsedSection] = [root]

    def current() -> ParsedSection:
        return stack[-1]

    def append_text(text: str) -> None:
        if not text.strip():
            return
        section = current()
        section.text = f"{section.text}\n{text}".strip() if section.text else text.strip()

    title: str | None = None

    for block in _iter_block_items(document):
        if isinstance(block, Table):
            append_text(_table_to_text(block))
            continue

        text = block.text.strip()
        if not text:
            continue

        level = _heading_level(block)
        if level is None:
            append_text(text)
            continue

        if level == 0:
            # The `Title` style names the document; it isn't a section of its
            # own (there's nothing to nest under it), so only the first one
            # is captured as metadata and none become tree nodes.
            if title is None:
                title = text
            continue

        while len(stack) > 1 and stack[-1].level >= level:
            stack.pop()

        new_section = ParsedSection(heading=text, level=level, text="", children=[])
        stack[-1].children.append(new_section)
        stack.append(new_section)

    sections = root.children
    if root.text.strip():
        preamble = ParsedSection(heading=None, level=0, text=root.text, children=[])
        sections = [preamble, *sections]

    return ParsedDocument(title=title, sections=sections)
