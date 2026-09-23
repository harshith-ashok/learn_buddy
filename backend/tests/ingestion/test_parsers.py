import io

import pytest
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter

from src.ingestion.parsers import UnsupportedDocumentFormatError, parse_document
from src.ingestion.parsers.docx import parse as parse_docx
from src.ingestion.parsers.pdf import parse as parse_pdf
from src.ingestion.parsers.pptx import parse as parse_pptx


def _build_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("Course Notes", style="Title")
    doc.add_paragraph("Some lead-in text before any section heading.")
    doc.add_heading("Chapter 1: Basics", level=1)
    doc.add_paragraph("Chapter 1 body text.")
    doc.add_heading("1.1 Details", level=2)
    doc.add_paragraph("Nested detail text.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "C"
    table.cell(1, 1).text = "D"
    doc.add_heading("Chapter 2: Advanced", level=1)
    doc.add_paragraph("Chapter 2 body text.")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _build_pptx_bytes() -> bytes:
    presentation = Presentation()
    title_slide_layout = presentation.slide_layouts[1]  # "Title and Content"

    slide1 = presentation.slides.add_slide(title_slide_layout)
    slide1.shapes.title.text = "Introduction"
    slide1.placeholders[1].text = "Welcome to the course."

    slide2 = presentation.slides.add_slide(title_slide_layout)
    slide2.shapes.title.text = "Topic Overview"
    slide2.placeholders[1].text = "Topic A\nTopic B"

    blank_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = blank_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text_frame.text = "Untitled slide body only"

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _build_pdf_bytes(page_texts: list[str]) -> bytes:
    writer = PdfWriter()
    for _ in page_texts:
        writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def test_docx_parser_builds_heading_tree_with_preamble_and_table() -> None:
    parsed = parse_docx(_build_docx_bytes())

    assert parsed.title == "Course Notes"
    headings = [s.heading for s in parsed.sections]
    assert headings == [None, "Chapter 1: Basics", "Chapter 2: Advanced"]

    preamble = parsed.sections[0]
    assert "lead-in text" in preamble.text

    chapter1 = parsed.sections[1]
    assert "Chapter 1 body text." in chapter1.text
    assert len(chapter1.children) == 1
    subsection = chapter1.children[0]
    assert subsection.heading == "1.1 Details"
    assert "Nested detail text." in subsection.text
    # the table (which follows "1.1 Details" in the source) nests there too,
    # with its row cells kept together
    assert "A | B" in subsection.text
    assert "C | D" in subsection.text

    chapter2 = parsed.sections[2]
    assert "Chapter 2 body text." in chapter2.text
    assert chapter2.children == []


def test_docx_parser_handles_no_headings() -> None:
    doc = DocxDocument()
    doc.add_paragraph("Just some plain notes with no headings at all.")
    buffer = io.BytesIO()
    doc.save(buffer)

    parsed = parse_docx(buffer.getvalue())

    assert len(parsed.sections) == 1
    assert parsed.sections[0].heading is None
    assert "plain notes" in parsed.sections[0].text


def test_pptx_parser_uses_slide_titles_as_headings() -> None:
    parsed = parse_pptx(_build_pptx_bytes())

    assert len(parsed.sections) == 3
    assert parsed.sections[0].heading == "Introduction"
    assert "Welcome to the course." in parsed.sections[0].text
    assert parsed.sections[1].heading == "Topic Overview"
    assert "Topic A" in parsed.sections[1].text
    assert parsed.sections[2].heading is None
    assert "Untitled slide body only" in parsed.sections[2].text
    assert parsed.title == "Introduction"


def test_pdf_parser_returns_empty_sections_for_blank_pages() -> None:
    parsed = parse_pdf(_build_pdf_bytes(["", ""]))
    assert parsed.sections == []


def test_dispatch_by_extension() -> None:
    parsed = parse_document("notes.docx", _build_docx_bytes())
    assert parsed.title == "Course Notes"


def test_dispatch_rejects_unsupported_extension() -> None:
    with pytest.raises(UnsupportedDocumentFormatError):
        parse_document("notes.txt", b"hello")
