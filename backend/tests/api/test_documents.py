import io
import uuid

from docx import Document as DocxDocument
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from src.core.llm_client import get_llm_client
from src.core.security import create_access_token, hash_password
from src.core.vectorstore import get_vectorstore_client
from src.db.models import Document, DocumentStatus, Student, Topic
from src.db.session import get_session_factory
from src.ingestion.schemas import ExtractedTopic, TopicExtractionResult
from src.main import app


class _ReuseSession:
    """An async-context-manager `session_factory()` call must return.

    Wraps an already-open session without closing it on `__aexit__`, so
    `ingestion.pipeline.process_document`'s `async with session_factory() as
    db:` reuses the test's transactional `db_session` instead of opening a
    connection to a different database.
    """

    def __init__(self, session: AsyncSession) -> None:
        self._session = session

    async def __aenter__(self) -> AsyncSession:
        return self._session

    async def __aexit__(self, *exc_info: object) -> None:
        return None


def _override_pipeline_dependencies(fake_llm_client, test_vectorstore, db_session: AsyncSession) -> None:
    app.dependency_overrides[get_llm_client] = lambda: fake_llm_client
    app.dependency_overrides[get_vectorstore_client] = lambda: test_vectorstore
    app.dependency_overrides[get_session_factory] = lambda: lambda: _ReuseSession(db_session)


def _build_syllabus_docx() -> bytes:
    """A small fixture syllabus: 2 chapters, each with enough body text for
    one chunk at the default ~400-word chunk size."""
    doc = DocxDocument()
    doc.add_paragraph("Intro to Testing", style="Title")
    doc.add_heading("Chapter 1: Unit Testing", level=1)
    doc.add_paragraph(" ".join(f"unit{i}" for i in range(60)))
    doc.add_heading("Chapter 2: Integration Testing", level=1)
    doc.add_paragraph(" ".join(f"integration{i}" for i in range(60)))

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


async def _make_authenticated_student(db_session: AsyncSession) -> tuple[Student, str]:
    student = Student(email="ingest@example.com", hashed_password=hash_password("x"), full_name="I")
    db_session.add(student)
    await db_session.flush()
    token = create_access_token(str(student.id))
    return student, token


async def test_upload_runs_pipeline_and_produces_expected_topics_and_chunks(
    client: AsyncClient, db_session: AsyncSession, fake_llm_client, test_vectorstore
) -> None:
    _student, token = await _make_authenticated_student(db_session)
    fake_llm_client.extraction = TopicExtractionResult(
        topics=[
            ExtractedTopic(name="Unit Testing", description="Testing units in isolation"),
            ExtractedTopic(
                name="Integration Testing",
                description="Testing components together",
                prerequisites=["Unit Testing"],
            ),
        ]
    )

    _override_pipeline_dependencies(fake_llm_client, test_vectorstore, db_session)

    response = await client.post(
        "/documents/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={
            "file": (
                "syllabus.docx",
                _build_syllabus_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["created"] is True

    document = await db_session.get(Document, uuid.UUID(body["id"]))
    assert document.status == DocumentStatus.DONE

    topics = (await db_session.scalars(select(Topic).where(Topic.document_id == document.id))).all()
    assert len(topics) == 2  # expected topic count, from the fake extraction result

    collection = test_vectorstore.get_or_create_collection(str(document.student_id))
    assert collection.count() == 2  # expected chunk count: one chunk per chapter


async def test_reupload_same_bytes_reuses_document_and_replaces_topics(
    client: AsyncClient, db_session: AsyncSession, fake_llm_client, test_vectorstore
) -> None:
    _student, token = await _make_authenticated_student(db_session)
    fake_llm_client.extraction = TopicExtractionResult(
        topics=[ExtractedTopic(name="Unit Testing", description="Testing units in isolation")]
    )
    _override_pipeline_dependencies(fake_llm_client, test_vectorstore, db_session)

    content = _build_syllabus_docx()
    file_field = (
        "syllabus.docx",
        content,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    first = await client.post(
        "/documents/upload", headers={"Authorization": f"Bearer {token}"}, files={"file": file_field}
    )
    second = await client.post(
        "/documents/upload", headers={"Authorization": f"Bearer {token}"}, files={"file": file_field}
    )

    assert first.json()["created"] is True
    assert second.json()["created"] is False
    assert first.json()["id"] == second.json()["id"]

    documents = (await db_session.scalars(select(Document))).all()
    assert len(documents) == 1

    topics = (await db_session.scalars(select(Topic).where(Topic.document_id == documents[0].id))).all()
    assert len(topics) == 1  # re-processed, not duplicated


async def test_upload_rejects_unsupported_extension(client: AsyncClient, db_session: AsyncSession) -> None:
    _student, token = await _make_authenticated_student(db_session)

    response = await client.post(
        "/documents/upload",
        headers={"Authorization": f"Bearer {token}"},
        files={"file": ("notes.txt", b"plain text", "text/plain")},
    )

    assert response.status_code == 415
    assert response.json()["code"] == "unsupported_media_type"


async def test_upload_requires_authentication(client: AsyncClient) -> None:
    response = await client.post(
        "/documents/upload",
        files={
            "file": (
                "syllabus.docx",
                _build_syllabus_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 401


async def test_list_documents_returns_only_the_caller_own_documents(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    _owner, owner_token = await _make_authenticated_student(db_session)
    other = Student(email="other-docs@example.com", hashed_password=hash_password("x"), full_name="O")
    db_session.add(other)
    await db_session.flush()

    mine = Document(
        student_id=_owner.id,
        filename="mine.pdf",
        storage_path="/tmp/mine.pdf",
        content_hash="mine-hash",
        status=DocumentStatus.DONE,
    )
    theirs = Document(
        student_id=other.id,
        filename="theirs.pdf",
        storage_path="/tmp/theirs.pdf",
        content_hash="theirs-hash",
        status=DocumentStatus.DONE,
    )
    db_session.add_all([mine, theirs])
    await db_session.flush()

    response = await client.get("/documents", headers={"Authorization": f"Bearer {owner_token}"})

    assert response.status_code == 200
    filenames = [doc["filename"] for doc in response.json()]
    assert filenames == ["mine.pdf"]


async def test_get_document_by_id(client: AsyncClient, db_session: AsyncSession) -> None:
    _owner, token = await _make_authenticated_student(db_session)
    document = Document(
        student_id=_owner.id,
        filename="mine.pdf",
        storage_path="/tmp/mine.pdf",
        content_hash="mine-hash",
        status=DocumentStatus.DONE,
    )
    db_session.add(document)
    await db_session.flush()

    response = await client.get(f"/documents/{document.id}", headers={"Authorization": f"Bearer {token}"})

    assert response.status_code == 200
    assert response.json()["filename"] == "mine.pdf"


async def test_get_document_for_a_foreign_document_is_not_found(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    _owner, token = await _make_authenticated_student(db_session)

    response = await client.get(f"/documents/{uuid.uuid4()}", headers={"Authorization": f"Bearer {token}"})

    assert response.status_code == 404
